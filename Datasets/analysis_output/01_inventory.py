import openpyxl, zipfile, os, json
from openpyxl.utils import get_column_letter

BASE="/home/bhavith/Documents/Datasets"
FILES=["APACHE 4.xlsx","NIMS TRIAGE 2023.xlsx","NIMS TRIAGE 2024 (Responses).xlsx"]

for f in FILES:
    p=os.path.join(BASE,f)
    print("="*100)
    print("FILE:",f,"| size:",os.path.getsize(p),"bytes")
    z=zipfile.ZipFile(p)
    names=z.namelist()
    print("  zip parts:",len(names))
    print("  parts sample:",[n for n in names if not n.startswith('xl/worksheets')][:25])
    # core props
    try:
        import re
        core=z.read('docProps/core.xml').decode('utf8',errors='replace')
        print("  CORE PROPS:",re.sub(r'\s+',' ',core)[:900])
    except Exception as e: print("  no core props",e)
    try:
        app=z.read('docProps/app.xml').decode('utf8',errors='replace')
        print("  APP PROPS:",re.sub(r'\s+',' ',app)[:900])
    except Exception as e: print("  no app props",e)

    wb=openpyxl.load_workbook(p, read_only=False, data_only=False)
    print("  SHEETS:",wb.sheetnames)
    for ws in wb.worksheets:
        print(f"  --- sheet '{ws.title}' state={ws.sheet_state} dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
        print(f"      merged={len(ws.merged_cells.ranges)} {[str(r) for r in list(ws.merged_cells.ranges)[:10]]}")
        print(f"      freeze={ws.freeze_panes} autofilter={ws.auto_filter.ref} tables={list(getattr(ws,'tables',{}) or {})}")
        # count formulas
        nf=0; ex=[]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value,str) and c.value.startswith('='):
                    nf+=1
                    if len(ex)<5: ex.append((c.coordinate,c.value))
        print(f"      formula_cells={nf} examples={ex}")
        # first 6 rows raw
        for i,row in enumerate(ws.iter_rows(min_row=1,max_row=6,values_only=True),1):
            vals=[('' if v is None else str(v))[:30] for v in row[:18]]
            print(f"      R{i}:",vals)
    wb.close()
