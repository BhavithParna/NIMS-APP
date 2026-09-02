import openpyxl, os
from openpyxl.utils import get_column_letter
BASE="/home/bhavith/Documents/Datasets"
SPEC=[("APACHE 4.xlsx",["Sheet1","Sheet2"]),
      ("NIMS TRIAGE 2023.xlsx",["Form Responses 1","Sheet2","Sheet3"]),
      ("NIMS TRIAGE 2024 (Responses).xlsx",["Form Responses 1"])]
for f,sheets in SPEC:
    wb=openpyxl.load_workbook(os.path.join(BASE,f),read_only=True,data_only=True)
    for s in sheets:
        ws=wb[s]
        print("="*90); print(f"{f} :: {s}  ({ws.max_row}r x {ws.max_column}c)")
        rows=list(ws.iter_rows(min_row=1,max_row=3,values_only=True))
        hdr=rows[0]
        for i,h in enumerate(hdr):
            L=get_column_letter(i+1)
            v1=rows[1][i] if len(rows)>1 and i<len(rows[1]) else None
            v2=rows[2][i] if len(rows)>2 and i<len(rows[2]) else None
            print(f"  {L:>3} [{i:>2}] {repr(h)[:60]:<62} ex1={repr(v1)[:28]:<30} ex2={repr(v2)[:28]}")
    wb.close()
